"""Raporu düzenlenebilir Excel (.xlsx) dosyasına yazar (openpyxl).

Sayılar Excel'e sayı olarak yazılır (Excel kendi yerel ayarıyla gösterir); metin olarak değil.
"""

from __future__ import annotations

import os
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from yapimmaliyet.branding import letterhead_path
from yapimmaliyet.reports.model import ColKind, ReportDoc

NUMBER_FORMATS = {ColKind.MONEY: "#,##0.00", ColKind.PRICE: "#,##0.00##", ColKind.QTY: "#,##0.000",
                  ColKind.PERCENT: "0.00%", ColKind.INT: "0"}
HEADER_FILL = PatternFill("solid", fgColor="ECE2E2")
THIN = Side(style="thin", color="CBBCBC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cell_value(v, kind: ColKind):
    if v is None or v == "":
        return None
    if kind is ColKind.TEXT:
        return str(v)
    if isinstance(v, Decimal):
        return v
    try:
        return Decimal(str(v))
    except ArithmeticError:
        return str(v)


LOGO_PX = 46   # Excel antetindeki logo yüksekliği (piksel)


def add_letterhead(ws, row: int = 1) -> int:
    """Sayfanın en üstüne antet (simge + logo). Logo yoksa satır eklemez. Döndürür: içeriğin ilk satırı."""
    logo = letterhead_path()
    if logo is None:
        return row
    img = XLImage(str(logo))
    img.width, img.height = round(LOGO_PX * img.width / img.height), LOGO_PX
    ws.add_image(img, f"A{row}")
    ws.row_dimensions[row].height = LOGO_PX * 0.75 + 8
    return row + 1


def write_xlsx(doc: ReportDoc, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"
    title = ws.cell(add_letterhead(ws), 1, doc.title + ("  — TASLAK (kesinleşmemiş)" if doc.draft else ""))
    title.font = Font(bold=True, size=14, color="A40000" if doc.draft else "000000")
    for k, v in doc.meta:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    widths: dict[int, float] = {}
    table_count = sum(section.table is not None for section in doc.sections)
    for section in doc.sections:
        ws.append([])
        ws.append([section.title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        for p in section.paragraphs:
            ws.append([p])
        t = section.table
        if t is None:
            continue
        ws.append([c.title for c in t.columns])
        header_row = ws.max_row
        for i, c in enumerate(t.columns, start=1):
            cell = ws.cell(header_row, i)
            cell.font, cell.fill, cell.border = Font(bold=True), HEADER_FILL, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            widths[i] = max(widths.get(i, 8), min(60, 8 + 10 * c.width))
        for row in t.rows + ([t.total_row] if t.total_row else []):
            ws.append([_cell_value(v, c.kind) for v, c in zip(row, t.columns)])
            r = ws.max_row
            for i, c in enumerate(t.columns, start=1):
                cell = ws.cell(r, i)
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="left" if c.kind is ColKind.TEXT else "right",
                                           vertical="top", wrap_text=c.kind is ColKind.TEXT)
                if c.kind in NUMBER_FORMATS:
                    cell.number_format = NUMBER_FORMATS[c.kind]
                if t.total_row is row:
                    cell.font = Font(bold=True)
        if not ws.freeze_panes or ws.freeze_panes == "A1":
            ws.freeze_panes = ws.cell(header_row + 1, 1)
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(t.columns))}{header_row + len(t.rows)}"
            if table_count == 1:
                ws.print_title_rows = f"{header_row}:{header_row}"
    if doc.notes:
        ws.append([])
        ws.append(["Notlar"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for n in doc.notes:
            ws.append([n])
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 14)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # report text must not become an executable Excel formula
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if doc.landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:" + ws.calculate_dimension().split(":")[-1]   # antet satırı da yazdırılır
    tmp = path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    os.replace(tmp, path)
    return path
