"""Belge dosyasını (Packet) Excel'e yazar: her belge ayrı çalışma sayfası, A4 sayfa yapısı ve imza alanlarıyla.

Sayılar Excel'e sayı olarak yazılır; sütun genişlikleri tabloların göreli genişliklerinden, yazdırma alanı ve
"sayfaya sığdır (genişlik)" ayarı her sayfa için kurulur. Taslak dosyada üst bilgi ve başlık TASLAK işaretlidir.
"""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from yapimmaliyet.reports.excel_writer import NUMBER_FORMATS, _cell_value, add_letterhead
from yapimmaliyet.reports.model import ColKind
from yapimmaliyet.reports.packet import KeyValues, Packet, Page, Signatures, TableBlock, Text

HEADER_FILL = PatternFill("solid", fgColor="ECE2E2")
BOLD_FILL = PatternFill("solid", fgColor="F4F0F0")
THIN = Side(style="thin", color="A39494")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RED = "A40000"
WRAP = Alignment(wrap_text=True, vertical="top")


def _safe_sheet(name: str, used: set[str]) -> str:
    base = "".join("-" if ch in '[]:*?/\\' else ch for ch in name)[:31] or "Sayfa"
    out, i = base, 2
    while out in used:
        out = f"{base[:28]}({i})"
        i += 1
    used.add(out)
    return out


def _ncols(page: Page) -> int:
    widths = [len(b.table.columns) for b in page.blocks if isinstance(b, TableBlock)]
    return max(widths + [4])


def _merge(ws: Worksheet, row: int, c1: int, c2: int) -> None:
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)


def _write_page(ws: Worksheet, page: Page, packet: Packet) -> None:
    n = _ncols(page)
    widths = {i: 12.0 for i in range(1, n + 1)}
    top = add_letterhead(ws)
    ws.cell(top, 1, packet.subtitle).font = Font(bold=True, size=11)
    _merge(ws, top, 1, n)
    title = page.title + ("  — TASLAK" if packet.draft else "")
    ws.cell(top + 1, 1, title).font = Font(bold=True, size=14, color=RED if packet.draft else "000000")
    _merge(ws, top + 1, 1, n)
    row = top + 2
    for k, v in page.header:
        ws.cell(row, max(1, n - 1), k).font = Font(bold=True)
        ws.cell(row, n, v).alignment = Alignment(horizontal="right")
        row += 1
    row += 1
    for block in page.blocks:
        if isinstance(block, Text):
            cell = ws.cell(row, 1, block.text)
            cell.font = Font(bold=block.style == "heading", italic=block.style == "note",
                             color=RED if block.style == "warning" else "000000", size=12 if block.style == "heading"
                             else 10)
            cell.alignment = WRAP
            _merge(ws, row, 1, n)
            ws.row_dimensions[row].height = max(15, 15 * (1 + len(block.text) // (12 * n)))
            row += 1
        elif isinstance(block, KeyValues):
            for k, v in block.rows:
                ws.cell(row, 1, k).font = Font(bold=True)
                ws.cell(row, 2, v).alignment = WRAP
                _merge(ws, row, 2, n)
                row += 1
            widths[1] = max(widths[1], 34)
            row += 1
        elif isinstance(block, TableBlock):
            row = _write_table(ws, block, row, widths)
        elif isinstance(block, Signatures):
            row = _write_signatures(ws, block, row + 1, n)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape" if page.landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(n)}{max(row, 1)}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = packet.footer_note[:200]
    ws.oddFooter.left.size = 7
    ws.oddFooter.right.text = "Sayfa &P / &N"
    if packet.draft:
        ws.oddHeader.center.text = "TASLAK — kontrol amaçlıdır"
        ws.oddHeader.center.color = RED


def _write_table(ws: Worksheet, block: TableBlock, row: int, widths: dict[int, float]) -> int:
    t = block.table
    if block.title:
        ws.cell(row, 1, block.title).font = Font(bold=True, size=11)
        row += 1
    for i, c in enumerate(t.columns, start=1):
        cell = ws.cell(row, i, c.title)
        cell.font, cell.fill, cell.border = Font(bold=True, size=9), HEADER_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        widths[i] = max(widths.get(i, 8), min(70, 5 + 9 * c.width))
    ws.row_dimensions[row].height = 30
    row += 1
    body = t.rows + ([t.total_row] if t.total_row else [])
    for r_i, values in enumerate(body):
        bold = r_i in block.bold_rows or (t.total_row is not None and r_i == len(body) - 1)
        for i, (v, c) in enumerate(zip(values, t.columns), start=1):
            cell = ws.cell(row, i, _cell_value(v, c.kind))
            cell.border = BORDER
            cell.font = Font(bold=bold, size=9)
            if bold:
                cell.fill = BOLD_FILL
            if c.kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[c.kind]
            if c.kind is ColKind.QTY:
                cell.number_format = "#,##0.000###"
            if c.kind is ColKind.TEXT:
                cell.alignment = WRAP
        row += 1
    return row + 1


def _write_signatures(ws: Worksheet, block: Signatures, row: int, n: int) -> int:
    items = block.items
    per_row = max(1, min(len(items), n))
    for start in range(0, len(items), per_row):
        chunk = items[start:start + per_row]
        seg = max(1, n // len(chunk))
        for j, sig in enumerate(chunk):
            c1 = 1 + j * seg
            c2 = n if j == len(chunk) - 1 else c1 + seg - 1
            for k, (text, font) in enumerate(((sig.role, Font(bold=True, size=9)), (sig.name, Font(size=9)),
                                              (sig.title, Font(italic=True, size=8)), ("İmza", Font(size=8,
                                                                                                    color="808080")))):
                cell = ws.cell(row + k, c1, text)
                cell.font, cell.alignment = font, Alignment(horizontal="center", wrap_text=True)
                _merge(ws, row + k, c1, c2)
            ws.cell(row + 3, c1).border = Border(top=THIN)
        if block.date_text:
            ws.cell(row + 4, 1, block.date_text).font = Font(size=8)
        row += 6
    return row


def write_packet_xlsx(packet: Packet, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for page in packet.pages:
        _write_page(wb.create_sheet(_safe_sheet(page.sheet, used)), page, packet)
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    wb.properties.title = packet.title
    tmp = path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    os.replace(tmp, path)
    return path
