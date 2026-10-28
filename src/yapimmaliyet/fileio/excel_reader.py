"""Excel (.xlsx) okuyucu.

- Bütün görünür sayfalar ayrı tablo olarak döner; kullanıcı hangisini kullanacağını seçer.
- Sayısal hücreler doğrudan Decimal olur (yerel ayara bağlı metin ayrıştırması gerekmez).
- Formülü olup hesaplanmış değeri kaydedilmemiş hücreler (dosya Excel dışında üretilmişse)
  sessizce boş sayılmaz; uyarı olarak bildirilir. Formüller yeniden hesaplanmaz: dosyada kayıtlı son sonuç
  okunur ve bu da bildirilir.
- Gizli sayfalar okunmaz (bildirilir). Gizli satırlar işaretlenir (önizlemede varsayılan olarak hariç), gizli
  sütunlar otomatik eşlemede kullanılmaz; ikisi de uyarıyla gösterilir.
- Dikey birleştirilmiş hücrelerde üst hücrenin değeri alt satırlara `merged` olarak taşınır; yalnızca iş grubu /
  kategori gibi metin alanlarında kullanılır (miktar ve tutar asla çoğaltılmaz).
"""

from __future__ import annotations

import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from yapimmaliyet.fileio.raw_table import CellValue, RawRow, RawTable, SourceKind
from yapimmaliyet.infra.errors import ValidationError


def _convert(value) -> CellValue:
    if value is None:
        return None
    if isinstance(value, bool):
        return "DOĞRU" if value else "YANLIŞ"
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        # 15 anlamlı basamak: kayan nokta kalıntılarını (0.30000000000000004) temizler
        return Decimal(format(value, ".15g"))
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value).strip()
    return text or None


def _formula_cells(path: Path) -> dict[str, set[tuple[int, int]]]:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        result: dict[str, set[tuple[int, int]]] = {}
        for ws in wb.worksheets:
            cells = set()
            for r_idx, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True), start=1):
                for c_idx, v in enumerate(row, start=1):
                    if isinstance(v, str) and v.startswith("="):
                        cells.add((r_idx, c_idx))
            result[ws.title] = cells
        return result
    finally:
        wb.close()


def _merged_values(ws) -> dict[tuple[int, int], CellValue]:
    """Dikey birleştirilmiş aralıklarda (tek sütun ya da çok sütun) üst satırın değerinin alt satırlardaki karşılığı."""
    out: dict[tuple[int, int], CellValue] = {}
    for rng in ws.merged_cells.ranges:
        if rng.max_row == rng.min_row:
            continue                                   # yalnız yatay birleştirme: taşınacak değer yok
        top = _convert(ws.cell(rng.min_row, rng.min_col).value)
        if top is None:
            continue
        for r in range(rng.min_row + 1, rng.max_row + 1):
            out[(r, rng.min_col)] = top
    return out


def read_xlsx(path: Path) -> list[RawTable]:
    try:
        formulas = _formula_cells(path)
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValidationError("Excel dosyası açılamadı. Dosya bozuk ya da parola korumalı olabilir.",
                              detail=str(exc)) from exc
    tables = []
    hidden_sheets = []
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                hidden_sheets.append(ws.title)
                continue
            rows: list[RawRow] = []
            missing: list[str] = []
            sheet_formulas = formulas.get(ws.title, set())
            merged = _merged_values(ws)
            hidden_rows = {i for i, d in ws.row_dimensions.items() if d.hidden}
            hidden_cols = _hidden_columns(ws)
            for r_idx, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True), start=1):
                values = [_convert(v) for v in row]
                for c_idx, v in enumerate(values, start=1):
                    if v is None and (r_idx, c_idx) in sheet_formulas:
                        missing.append(f"satır {r_idx} sütun {c_idx}")
                while values and values[-1] is None:
                    values.pop()
                inherited = {c - 1: v for (r, c), v in merged.items() if r == r_idx} or None
                if values:
                    rows.append(RawRow(values, f"{ws.title} satır {r_idx}", line=r_idx,
                                       hidden=r_idx in hidden_rows, merged=inherited))
            if not rows:
                continue
            table = RawTable(ws.title, rows, SourceKind.EXCEL, hidden_columns=hidden_cols)
            if missing:
                table.warnings.append(
                    f"{len(missing)} formül hücresinin hesaplanmış değeri dosyada yok (ör. {missing[0]}). "
                    "Dosyayı Excel'de açıp kaydederek değerlerin hesaplanmasını sağlayın.")
            if len(sheet_formulas) > len(missing):
                table.warnings.append(
                    f"{len(sheet_formulas) - len(missing)} formül hücresi dosyada kayıtlı son hesaplanmış değeriyle "
                    "okundu; formüller yeniden hesaplanmaz. Dosya Excel dışında düzenlendiyse değerleri kontrol edin.")
            n_hidden = sum(1 for r in rows if r.hidden)
            if n_hidden:
                table.warnings.append(f"{n_hidden} satır Excel'de gizli; önizlemede işaretli ve varsayılan olarak "
                                      "aktarılmaz.")
            if hidden_cols:
                letters = ", ".join(get_column_letter(c + 1) for c in sorted(hidden_cols))
                table.warnings.append(f"Gizli sütunlar ({letters}) otomatik eşlemede kullanılmadı.")
            if merged:
                table.warnings.append(f"{len(merged)} hücre dikey birleştirilmiş aralıkta; değer yalnızca iş grubu / "
                                      "kategori alanında alt satırlara taşınır.")
            tables.append(table)
    finally:
        wb.close()
    if not tables:
        raise ValidationError("Excel dosyasında veri içeren görünür sayfa bulunamadı."
                              + (f" Gizli sayfalar: {', '.join(hidden_sheets)}." if hidden_sheets else ""))
    if hidden_sheets:
        tables[0].warnings.append(f"Gizli sayfalar okunmadı: {', '.join(hidden_sheets)}.")
    return tables


def _hidden_columns(ws) -> set[int]:
    from openpyxl.utils import column_index_from_string

    out: set[int] = set()
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            first = column_index_from_string(key)
            last = max(first, dim.max or first)
            out.update(range(first - 1, last))
    return out
