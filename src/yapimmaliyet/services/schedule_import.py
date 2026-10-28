"""İş programı (faaliyet listesi) dosyasından içe aktarma: Excel, CSV ve metin tabanlı PDF.

Belirli bir yazılımın biçimine bağlı değildir: bir başlık satırı ve altında faaliyetler aranır. Tanınan başlıklar
(Türkçe ve İngilizce, ör. MS Project/Primavera dışa aktarımları): Faaliyet / İş / Görev / Task Name, Başlangıç / Start,
Bitiş / Finish, Planlanan maliyet / Maliyet / Cost, İlerleme / % Complete, Disiplin, Birim, Miktar, Not.
Tarihler GG.AA.YYYY, GG/AA/YYYY ve YYYY-AA-GG kabul edilir. İlerleme 0–1 aralığındaysa ve sütundaki bütün değerler ≤ 1 ise
kesir (×100), aksi hâlde yüzde sayılır. Hiçbir satır sessizce atılmaz; reddedilenler nedenleriyle listelenir ve fark varsa
açık onay olmadan aktarılmaz. Dosyada maliyet yoksa 0 yazılır ve uyarılır; plan maliyeti sonradan girilir.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from yapimmaliyet.domain.parsing import parse_user_decimal
from yapimmaliyet.domain.tender import DISCIPLINES
from yapimmaliyet.domain.text_tr import normalize_search
from yapimmaliyet.fileio import csv_reader, excel_reader, pdf_text
from yapimmaliyet.fileio.detect import FileKind, detect_kind
from yapimmaliyet.infra.errors import ValidationError

HEADERS = {
    "code": ("aktivite kodu", "aktivite id", "faaliyet kodu", "kod", "id", "activity id", "external code", "task id", "aktivite kodu id", "kod id"),
    "wbs": ("wbs", "wbs kodu", "is kirilimi", "wbs code"),
    "name": ("is tanimi", "faaliyet tanimi", "tanim", "faaliyet", "faaliyet adi", "is", "is adi", "gorev", "gorev adi", "task name", "name", "aktivite", "activity"),
    "start": ("baslangic", "baslangic tarihi", "planlanan baslangic", "start", "start date"),
    "end": ("bitis", "bitis tarihi", "planlanan bitis", "finish", "end", "end date"),
    "cost": ("plan tutari", "planlanan tutar", "butce tutari", "planlanan maliyet", "maliyet", "butce", "cost", "tutar"),
    "progress": ("ilerleme", "ilerleme %", "tamamlanma", "% complete", "yuzde", "gerceklesme"),
    "discipline": ("disiplin", "discipline"),
    "unit": ("birim",),
    "qty": ("miktar", "planlanan miktar"),
    "notes": ("not", "aciklama", "notes"),
}


@dataclass
class ActivityRow:
    ref: str
    data: dict


@dataclass
class Reject:
    ref: str
    text: str
    reason: str


@dataclass
class ScheduleImport:
    sheet: str
    source_rows: int = 0
    rows: list[ActivityRow] = field(default_factory=list)            # eklenecek yeni faaliyetler
    updates: list[tuple[int, ActivityRow]] = field(default_factory=list)   # güncellenecek mevcut faaliyetler (kod eşleşmesi)
    skipped: list[ActivityRow] = field(default_factory=list)         # projede zaten kayıtlı (atlanır)
    rejected: list[Reject] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def clean(self) -> bool:
        return not self.rejected

    def summary(self) -> str:
        return (f"'{self.sheet}': {self.source_rows} satır okundu, {len(self.rows)} yeni faaliyet, {len(self.updates)} güncelleme, "
                f"{len(self.skipped)} zaten kayıtlı (atlanır), {len(self.rejected)} satır reddedildi.")


def _key(v) -> str:
    return normalize_search(str(v or "")).strip()


def _header(table):
    for i, row in enumerate(table.rows[:40]):
        cols: dict[str, int] = {}
        for j, v in enumerate(row.values):
            raw = str(v or "").strip().lower()
            k = _key(v)
            if not k and raw not in ("%",):
                continue
            for name, words in HEADERS.items():
                if name not in cols and (k in words or raw in words):
                    cols[name] = j
        if {"name", "start", "end"} <= cols.keys():
            return i, cols
    return None


def _date(v) -> str:
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d.%m.%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Tarih okunamadı: '{s}'")


def _dec(v) -> Decimal | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, Decimal):
        return v
    s = re.sub(r"[%₺TLtl\s]", "", str(v))
    return parse_user_decimal(s)


def _disc(v) -> str | None:
    k = _key(v)
    if not k:
        return None
    for code, label in DISCIPLINES.items():
        if k in (normalize_search(code), normalize_search(label)):
            return code
    raise ValueError(f"Disiplin tanınmadı: '{v}'")


def _getter(row, cols: dict[str, int]):
    """Satıra bağlı hücre okuyucu (döngü değişkenine geç bağlanmaz)."""
    return lambda n: row.cell(cols[n]) if n in cols else None


def _load_tables(path: Path):
    kind = detect_kind(path)
    try:
        if kind is FileKind.XLSX:
            tables = excel_reader.read_xlsx(path)
        elif kind is FileKind.CSV:
            tables = [csv_reader.read_csv(path)]
        else:
            pdf = pdf_text.analyze_pdf(path)
            if not pdf.text_pages:
                raise ValidationError("Bu PDF taranmış görüntü içeriyor; metin okunamadı. Metin tabanlı PDF ya da Excel kullanın.")
            tables = [pdf_text.read_pdf_text(path, pdf.text_pages)]
    except ValidationError:
        raise
    except Exception as exc:   # bozuk dosya anlaşılır hata olarak döner
        raise ValidationError("İş programı dosyası okunamadı (bozuk ya da desteklenmeyen biçim).", detail=str(exc)) from exc
    return tables


def inspect_schedule_file(path: Path | str) -> dict:
    """Sütun eşleme ekranı için: ilk tablonun başlıkları, otomatik eşleme ve örnek satırlar (hiçbir şey kaydetmez)."""
    tables = _load_tables(Path(path))
    for t in tables:
        h = _header(t)
        if h:
            hrow, cols = h
            break
    else:
        t, hrow, cols = tables[0], 0, {}
    heads = [str(v or "") for v in t.rows[hrow].values] if t.rows else []
    sample = [[str(v or "") for v in r.values] for r in t.rows[hrow + 1:hrow + 6]]
    return {"sheet": t.name, "header_row": hrow, "headers": heads, "auto_mapping": cols, "sample": sample,
            "fields": {"code": "Faaliyet kodu", "wbs": "WBS kodu", "name": "Faaliyet adı *", "start": "Başlangıç *",
                       "end": "Bitiş *", "cost": "Planlanan maliyet", "progress": "İlerleme %", "discipline": "Disiplin",
                       "unit": "Birim", "qty": "Miktar", "notes": "Not"}}


def read_schedule_file(path: Path | str, mapping: dict[str, int] | None = None) -> ScheduleImport:
    """`mapping`: kanonik alan → sütun indeksi; verilirse otomatik başlık eşlemesinin yerine geçer (zorunlu: name, start, end)."""
    path = Path(path)
    tables = _load_tables(path)
    found = [(t, _header(t)) for t in tables]
    found = [(t, h) for t, h in found if h]
    if mapping:
        if not {"name", "start", "end"} <= mapping.keys():
            raise ValidationError("Sütun eşlemesinde Faaliyet adı, Başlangıç ve Bitiş zorunludur.")
        base = found[0] if found else (tables[0], (0, {}))
        table, hrow = base[0], base[1][0]
        cols = {k: int(v) for k, v in mapping.items() if v is not None}
    elif found:
        table, (hrow, cols) = found[0]
    else:
        raise ValidationError("Dosyada iş programı tablosu bulunamadı: 'Faaliyet', 'Başlangıç' ve 'Bitiş' sütunlarını içeren "
                              "bir başlık satırı gerekir (ör. Task Name / Start / Finish) ya da sütunları elle eşleyin. "
                              "Şablonu kaydedip doldurabilirsiniz.")
    res = ScheduleImport(table.name)
    if len(found) > 1 and not mapping:
        res.warnings.append(f"{len(found)} sayfada iş programı var; yalnız '{table.name}' okundu.")
    raw_rows = []
    for row in table.rows[hrow + 1:]:
        get = _getter(row, cols)
        name = str(get("name") or "").strip()
        if not name and not str(get("start") or "").strip() and not str(get("cost") or "").strip():
            continue
        raw_rows.append((row.ref, name, get))
    res.source_rows = len(raw_rows)
    progress_vals = []
    for _ref, _n, get in raw_rows:
        try:
            v = _dec(get("progress"))
        except ValueError:
            v = None
        if v is not None:
            progress_vals.append(v)
    as_fraction = bool(progress_vals) and max(progress_vals) <= 1 and "progress" in cols
    if as_fraction:
        res.warnings.append("İlerleme sütunundaki değerlerin hepsi 1 veya altında: kesir olarak (×100) yorumlandı.")
    seen: set[tuple] = set()
    for ref, name, get in raw_rows:
        try:
            if not name:
                raise ValueError("Faaliyet adı boş")
            start, end = _date(get("start")), _date(get("end"))
            if end < start:
                raise ValueError("Bitiş tarihi başlangıçtan önce")
            cost = _dec(get("cost"))
            prog = _dec(get("progress")) or Decimal(0)
            prog = prog * 100 if as_fraction else prog
            if not Decimal(0) <= prog <= Decimal(100):
                raise ValueError("İlerleme 0–100 arasında olmalı")
            if cost is not None and cost < 0:
                raise ValueError("Maliyet negatif")
            disc = _disc(get("discipline"))
            qty = _dec(get("qty"))
        except ValueError as exc:
            res.rejected.append(Reject(ref, name, str(exc)))
            continue
        code = str(get("code") or "").strip() or None
        key = (code,) if code else (_key(name), start, end)
        if key in seen:
            res.rejected.append(Reject(ref, name, "Aynı faaliyet kodu dosyada yinelenmiş" if code else
                                       "Aynı faaliyet (ad ve tarihler) dosyada yinelenmiş"))
            continue
        seen.add(key)
        if cost is None:
            res.warnings.append(f"{ref}: planlanan maliyet yok, 0 olarak alındı.")
        res.rows.append(ActivityRow(ref, dict(external_code=code, wbs_code=str(get("wbs") or "").strip() or None,
                                              name=name, discipline=disc, planned_start=start, planned_end=end,
                                              planned_cost=cost if cost is not None else Decimal(0), progress=prog,
                                              planned_qty=qty, unit=str(get("unit") or "").strip() or None,
                                              notes=str(get("notes") or "").strip() or None)))
    return res


def write_template(path: Path | str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Programı"
    heads = ["Faaliyet", "Disiplin", "Başlangıç", "Bitiş", "Planlanan maliyet", "İlerleme %", "Birim", "Miktar", "Not"]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(["Örnek faaliyet", "İnşaat", date(2027, 3, 1), date(2027, 5, 31), 1000000, 0, "", None, ""])
    for col, w in zip("ABCDEFGHI", (36, 12, 14, 14, 18, 12, 10, 12, 30)):
        ws.column_dimensions[col].width = w
    path = Path(path)
    wb.save(path)
    return path
