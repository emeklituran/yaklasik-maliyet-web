"""Stok girişi için Excel okuma / şablon yazma. Dosya yalnız satırları verir; kaydı StockService.receive yapar.

Beklenen sütunlar (başlık adı Türkçe karakter/büyük-küçük harf farkından bağımsız aranır): Kalem (açıklama, malzeme, ürün),
Birim, Miktar (adet) ve isteğe bağlı Poz No. Başlık satırı ilk 15 satır içinde aranır; boş satırlar atlanır.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook

from yapimmaliyet.domain.parsing import parse_user_decimal
from yapimmaliyet.domain.text_tr import normalize_search
from yapimmaliyet.infra.errors import ValidationError

NAMES = {"description": ("kalem", "aciklama", "malzeme", "urun", "tanim", "imalat"), "unit": ("birim",),
         "quantity": ("miktar", "adet", "teslim miktari"), "code": ("poz no", "poz", "kod", "malzeme kodu")}
TEMPLATE_HEAD = ["Poz No", "Kalem", "Birim", "Miktar"]


@dataclass
class StockSheet:
    rows: list[dict] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


def _norm(v) -> str:
    return normalize_search(str(v or "")).strip()


def read_stock_excel(path: str | Path) -> StockSheet:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("Excel dosyası açılamadı.", detail=str(exc)) from exc
    ws = wb.active
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    head_at, cols = None, {}
    for i, row in enumerate(grid[:15]):
        found = {}
        for j, cell in enumerate(row):
            n = _norm(cell)
            for key, names in NAMES.items():
                if key not in found and n in names:
                    found[key] = j
        if {"description", "quantity"} <= set(found):
            head_at, cols = i, found
            break
    if head_at is None:
        raise ValidationError("Başlık satırı bulunamadı. 'Kalem', 'Birim' ve 'Miktar' sütunları gerekir "
                              "(şablon için 'Şablon İndir' düğmesini kullanın).")
    sheet = StockSheet()
    for i, row in enumerate(grid[head_at + 1:], start=head_at + 2):
        def cell(key):
            j = cols.get(key)
            return row[j] if j is not None and j < len(row) else None
        desc = str(cell("description") or "").strip()
        raw_qty = cell("quantity")
        if not desc and raw_qty in (None, ""):
            continue
        try:
            qty = raw_qty if isinstance(raw_qty, Decimal) else (
                Decimal(str(raw_qty)) if isinstance(raw_qty, (int, float)) else parse_user_decimal(str(raw_qty or "")))
        except (InvalidOperation, ValueError, ValidationError):
            sheet.problems.append(f"Satır {i}: miktar okunamadı ('{raw_qty}').")
            continue
        unit = str(cell("unit") or "").strip()
        if not desc or qty <= 0 or not unit:
            sheet.problems.append(f"Satır {i}: kalem, birim ve sıfırdan büyük miktar gerekir.")
            continue
        sheet.rows.append({"code": str(cell("code") or "").strip(), "description": desc, "unit": unit, "quantity": qty})
    return sheet


def write_template(path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stok girişi"
    ws.append(TEMPLATE_HEAD)
    ws.append(["15.150.1001", "Hazır beton C25/30", "m³", 12.5])
    for col, w in zip("ABCD", (16, 48, 10, 12)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return Path(path)
